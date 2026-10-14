# -*- coding: utf-8 -*-
"""把《考拉兹截断_完整证明.tex》转成 Word（.docx），供 LibreOffice 生成 PDF。

只支持这份文档用到的 LaTeX 子集；遇到不认识的命令不静默放过，会在末尾报告。
数学用 Unicode 近似排版（与既有 .docx/.pdf 的风格一致），不追求 TeX 级公式排版。
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "考拉兹截断_完整证明.tex")
DST = os.path.join(HERE, "考拉兹截断_完整证明.docx")

NAMES = {"theorem": "定理", "lemma": "引理", "proposition": "命题", "corollary": "推论",
         "definition": "定义", "conjecture": "猜想", "remark": "注"}
NUMBERED = list(NAMES)

# ---------------------------------------------------------------- 数学转换
MAP = {
    "\\approx": "≈", "\\bigcup": "∪", "\\bigcap": "∩", "\\cap": "∩", "\\cup": "∪",
    "\\cdot": "·", "\\cdots": "⋯", "\\div": "÷", "\\dots": "…", "\\equiv": "≡",
    "\\exists": "∃", "\\forall": "∀", "\\ge": "≥", "\\le": "≤", "\\ne": "≠",
    "\\in": "∈", "\\notin": "∉", "\\infty": "∞", "\\int": "∫", "\\kappa": "κ",
    "\\theta": "θ", "\\varepsilon": "ε", "\\lceil": "⌈", "\\rceil": "⌉",
    "\\lfloor": "⌊", "\\rfloor": "⌋", "\\liminf": "liminf", "\\limsup": "limsup",
    "\\ln": "ln", "\\log": "log", "\\min": "min", "\\max": "max", "\\gcd": "gcd",
    "\\longrightarrow": "⟶", "\\to": "→", "\\mapsto": "↦", "\\mathbb": "Z",
    "\\mid": "∣", "\\nmid": "∤", "\\not\\subseteq": "⊄", "\\not\\equiv": "≢",
    "\\not": "¬", "\\qquad": "  ", "\\quad": " ", "\\;": " ", "\\:": " ",
    "\\,": "", "\\!": "", "\\ ": " ", "\\%": "%", "\\#": "#",
    "\\sqcup": "⊔", "\\subseteq": "⊆", "\\subsetneq": "⊊", "\\supseteq": "⊇",
    "\\sum": "Σ", "\\times": "×", "\\pm": "±", "\\varnothing": "∅",
    "\\left": "", "\\right": "", "\\big": "", "\\Big": "", "\\bigl": "", "\\bigr": "",
    "\\Bigl": "", "\\Bigr": "", "\\bigg": "", "\\Bigg": "",
    "\\dfrac": "", "\\frac": "", "\\tfrac": "", "\\displaystyle": "", "\\text": "",
}
SUP = {"0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴", "5": "⁵", "6": "⁶", "7": "⁷",
       "8": "⁸", "9": "⁹", "k": "ᵏ", "n": "ⁿ", "d": "ᵈ", "s": "ˢ", "t": "ᵗ", "K": "ᴷ",
       "a": "ᵃ", "b": "ᵇ", "j": "ʲ", "m": "ᵐ", "i": "ⁱ", "+": "⁺", "-": "⁻", "−": "⁻"}
SUB = {"0": "₀", "1": "₁", "2": "₂", "3": "₃", "4": "₄", "5": "₅", "6": "₆", "7": "₇",
       "8": "₈", "9": "₉", "k": "ₖ", "n": "ₙ", "a": "ₐ", "i": "ᵢ", "j": "ⱼ", "m": "ₘ",
       "r": "ᵣ", "s": "ₛ", "t": "ₜ", "x": "ₓ"}


def _tex_arg(s, i):
    while i < len(s) and s[i] == " ":
        i += 1
    if i >= len(s):
        return None
    if s[i] == "{":
        d, j = 0, i
        while j < len(s):
            if s[j] == "{":
                d += 1
            elif s[j] == "}":
                d -= 1
                if d == 0:
                    break
            j += 1
        return s[i + 1:j], j + 1
    return s[i], i + 1


def _frac(s):
    for cmd in ("\\dfrac", "\\tfrac", "\\frac"):
        while cmd in s:
            k = s.index(cmd)
            a = _tex_arg(s, k + len(cmd))
            if a is None:
                s = s[:k]
                continue
            num, i2 = a
            b = _tex_arg(s, i2)
            if b is None:
                s = s[:k] + num
                continue
            den, i3 = b
            s = s[:k] + "(" + num + ")/(" + den + ")" + s[i3:]
    return s


def m2u(s):
    """数学片段 -> Unicode 近似。"""
    s = _frac(s)
    s = re.sub(r"\\pmod\s*\{([^}]*)\}", r"(mod \1)", s)
    s = re.sub(r"\\pmod\s*([0-9A-Za-z])", r"(mod \1)", s)
    s = re.sub(r"\\bmod\s*", " mod ", s)
    s = re.sub(r"\\xrightarrow\s*\{([^}]*)\}", lambda m: "→(" + m.group(1) + ")", s)
    s = re.sub(r"\\tilde\s*([A-Za-z])", lambda m: m.group(1) + "\u0303", s)
    s = re.sub(r"\\mathbb\s*\{([A-Za-z])\}", r"\1", s)
    s = re.sub(r"\\mathbb\s*([A-Za-z])", r"\1", s)
    s = re.sub(r"\\boxed\s*", "", s)
    s = s.replace("\\{", "\x01").replace("\\}", "\x02")
    for a, b in sorted(MAP.items(), key=lambda x: -len(x[0])):
        s = s.replace(a, b)
    out, i = [], 0
    while i < len(s):
        if s[i] == "^" and i + 1 < len(s):
            if s[i + 1] == "{":
                j = s.find("}", i)
                if j > 0:
                    out.append("^(" + s[i + 2:j] + ")"); i = j + 1; continue
            out.append(SUP.get(s[i + 1], "^" + s[i + 1])); i += 2; continue
        if s[i] == "_" and i + 1 < len(s):
            if s[i + 1] == "{":
                j = s.find("}", i)
                if j > 0:
                    out.append("_" + s[i + 2:j]); i = j + 1; continue
            out.append(SUB.get(s[i + 1], "_" + s[i + 1])); i += 2; continue
        out.append(s[i]); i += 1
    s = "".join(out).replace("\\", "").replace("{", "").replace("}", "")
    return s.replace("∤ ", "∤").replace("\x01", "{").replace("\x02", "}")


# ---------------------------------------------------------------- 解析
def match_brace(s, i):
    """s[i] == '{'，返回配对 '}' 的下标。"""
    d = 0
    for j in range(i, len(s)):
        if s[j] == "{":
            d += 1
        elif s[j] == "}":
            d -= 1
            if d == 0:
                return j
    return len(s) - 1


def blocks_of(txt):
    """把一段 LaTeX 文本切成 [('t',文字) | ('m',公式) | ('ul'|'ol',[条目])]。"""
    res, pos = [], 0
    pat = re.compile(r"\\\[(.*?)\\\]|\\begin\{(itemize|enumerate)\}(.*?)\\end\{\2\}", re.S)
    for m in pat.finditer(txt):
        if m.start() > pos:
            res.append(("t", txt[pos:m.start()]))
        if m.group(1) is not None:
            res.append(("m", m.group(1)))
        else:
            items = [x.strip() for x in re.split(r"\\item\b", m.group(3)) if x.strip()]
            res.append(("ul" if m.group(2) == "itemize" else "ol", items))
        pos = m.end()
    if pos < len(txt):
        res.append(("t", txt[pos:]))
    return res


def parse_body(tex):
    body = tex.split(r"\begin{document}", 1)[1].rsplit(r"\end{document}", 1)[0]
    body = "\n".join(l for l in body.split("\n") if not l.lstrip().startswith("%"))
    lines = body.split("\n")
    blocks, i = [], 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1; continue
        m = re.match(r"\\begin\{([A-Za-z*]+)\}(\[[^\]]*\])?", s)
        if m:
            env, opt = m.group(1), (m.group(2) or "")[1:-1]
            depth, collected, j = 0, [], i
            while j < len(lines):
                depth += len(re.findall(r"\\begin\{" + re.escape(env) + r"\}", lines[j]))
                depth -= len(re.findall(r"\\end\{" + re.escape(env) + r"\}", lines[j]))
                collected.append(lines[j]); j += 1
                if depth == 0:
                    break
            inner = "\n".join(collected)
            inner = re.sub(r"^\s*\\begin\{[^}]*\}(\[[^\]]*\])?", "", inner, count=1)
            inner = re.sub(r"\\end\{[^}]*\}\s*$", "", inner, count=1)
            blocks.append(("env", env, opt, inner)); i = j; continue
        for cmd, lvl in ((r"\section", 1), (r"\subsection", 2), (r"\subsubsection", 3)):
            mm = re.match(re.escape(cmd) + r"\{(.*)\}", s)
            if mm:
                blocks.append(("sec", lvl, mm.group(1))); break
        else:
            if s.startswith(r"\["):
                depth, collected, j = 0, [], i
                while j < len(lines):
                    collected.append(lines[j])
                    if r"\]" in lines[j] and j > i:
                        break
                    if j == i and r"\]" in lines[j][2:]:
                        break
                    j += 1
                inner = "\n".join(collected)
                inner = inner.strip()[2:]
                inner = inner[:inner.rfind(r"\]")] if r"\]" in inner else inner
                blocks.append(("env", "display", "", inner)); i = j + 1; continue
            if s.startswith(r"\bibitem"):
                blocks.append(("bib", s)); i += 1; continue
            if s.startswith(r"\maketitle") or s.startswith(r"\date") or s.startswith(r"\title") or s.startswith(r"\author"):
                i += 1; continue
            blocks.append(("p", s)); i += 1; continue
        i += 1
    return blocks


def number_labels(blocks):
    """复刻 LaTeX [section] 计数：返回 (label->编号, 块下标->编号)。"""
    lab2num, envnum = {}, {}
    sec, cnt = 0, 0
    for idx, b in enumerate(blocks):
        if b[0] == "sec" and b[1] == 1:
            sec += 1; cnt = 0
        elif b[0] == "env" and b[1] in NUMBERED:
            cnt += 1
            envnum[idx] = "%d.%d" % (sec, cnt)
            m = re.search(r"\\label\{([^}]+)\}", b[3])
            if m:
                lab2num[m.group(1)] = envnum[idx]
    return lab2num, envnum


# ---------------------------------------------------------------- 写出
def eastasia(run, font="宋体"):
    rPr = run._element.get_or_add_rPr()
    rF = rPr.find(qn('w:rFonts'))
    if rF is None:
        rF = OxmlElement('w:rFonts'); rPr.append(rF)
    rF.set(qn('w:eastAsia'), font)


def setf(run, size=None, math=True):
    if size:
        run.font.size = size
    run.font.name = "Cambria Math" if math else "宋体"
    eastasia(run)


WARN = []


def join_cjk_lines(txt):
    """LaTeX 源码的换行只是排版换行；中文之间不留空格，西文之间留一个空格。"""
    cjk = lambda ch: ord(ch) > 0x2E80
    out = ""
    for line in (l.strip() for l in txt.split("\n")):
        if not line:
            continue
        if out:
            out += "" if (cjk(out[-1]) or cjk(line[0])) else " "
        out += line
    return out


def emit_inline(par, txt, bold=False, italic=False, size=None, cjk=None):
    """把一段行内 LaTeX 写进段落。"""
    txt = join_cjk_lines(txt)
    i, buf = 0, ""

    def flush():
        nonlocal buf
        if buf:
            r = par.add_run(buf); r.bold = bold; r.italic = italic; setf(r, size)
            if cjk:
                eastasia(r, cjk)
            buf = ""

    while i < len(txt):
        if txt.startswith(r"\\", i):
            flush(); i += 2; continue
        if txt.startswith(r"\(", i):
            j = txt.find(r"\)", i)
            j = len(txt) if j < 0 else j
            buf += m2u(txt[i + 2:j]); i = j + 2; continue
        if txt.startswith(r"\[", i):
            j = txt.find(r"\]", i)
            j = len(txt) if j < 0 else j
            buf += m2u(txt[i + 2:j]); i = j + 2; continue
        mm = re.match(r"\\(textbf|emph|textit|text|mathrm|boxed)\{", txt[i:])
        if mm:
            j = match_brace(txt, i + len(mm.group(0)) - 1)
            flush()
            inner = txt[i + len(mm.group(0)):j]
            if mm.group(1) in ("textbf",):
                emit_inline(par, inner, True, italic, size, cjk)
            elif mm.group(1) in ("emph", "textit"):
                emit_inline(par, inner, bold, True, size, cjk)
            else:
                buf += m2u(inner)
            i = j + 1; continue
        mm = re.match(r"\\ref\{", txt[i:])
        if mm:
            j = txt.find("}", i)
            buf += LABELS.get(txt[i + 5:j], "??"); i = j + 1; continue
        if txt[i] == "~":
            buf += " "; i += 1; continue
        if txt[i] in "{}":
            i += 1; continue
        if txt[i] == "\\":
            mm = re.match(r"\\[a-zA-Z]+", txt[i:])
            if mm:
                WARN.append(mm.group(0))
                i += len(mm.group(0)); continue
        buf += txt[i]; i += 1
    flush()


def add_par(doc, txt, indent=True, center=False, size=None, bold=False, cjk=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent and not center:
        p.paragraph_format.first_line_indent = Pt(21)
    emit_inline(p, txt, bold=bold, size=size, cjk=cjk)
    return p


def render_segments(doc, segs, size=None):
    for kind, val in segs:
        if kind == "m":
            render_display(doc, val)
        elif kind == "t":
            if val.strip():
                add_par(doc, val, size=size)
        else:
            style = "List Bullet" if kind == "ul" else "List Number"
            for it in val:
                p = doc.add_paragraph(style=style)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                emit_inline(p, it, size=size)


def flat(txt):
    """把文本里的行内公式转成 Unicode，去掉 \\( \\) 外壳，供标题等使用。"""
    txt = re.sub(r"\\\((.*?)\\\)", lambda m: m2u(m.group(1)), txt, flags=re.S)
    txt = re.sub(r"\\[a-zA-Z]+", "", txt)
    return txt.replace("~", " ")


def render_display(doc, math):
    """\[...\]：处理 cases / array / 普通多行公式。"""
    if r"\begin{cases}" in math:
        pre = math.split(r"\begin{cases}")[0].strip()
        inner = math.split(r"\begin{cases}")[1].split(r"\end{cases}")[0]
        rows = [r for r in inner.split(r"\\") if r.strip()]
        parts = []
        for row in rows:
            row = re.sub(r"^\s*\[\s*[\d.]*\s*(mm|pt|em|ex|cm)\s*\]", "", row)   # 去掉 \\[2mm]
            cells = [m2u(c).strip().rstrip("，,") for c in row.split("&") if c.strip()]
            parts.append("，".join(cells))
        add_par(doc, m2u(pre) + " " + "；".join(parts), indent=False, center=True)
        return
    if r"\begin{array}" in math:
        pre = math.split(r"\begin{array}")[0].strip()
        inner = math.split(r"\begin{array}")[1].split(r"\end{array}")[0]
        inner = inner[inner.find("}") + 1:]
        rows = [r.replace(r"\hline", "") for r in inner.split(r"\\") if r.strip()]
        cells = [[c.strip() for c in r.split("&")] for r in rows]
        if pre:
            add_par(doc, pre, indent=False, center=True)
        t = doc.add_table(rows=len(cells), cols=max(len(c) for c in cells))
        t.style = "Light Grid Accent 1"
        for ri, row in enumerate(cells):
            for ci, cell in enumerate(row):
                p = t.cell(ri, ci).paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                emit_inline(p, cell if "\\(" in cell else m2u(cell), size=Pt(9))
                if ri == 0:
                    for r in p.runs:
                        r.bold = True
        return
    for line in [x for x in math.split(r"\\") if x.strip()]:
        add_par(doc, m2u(line.strip()), indent=False, center=True)


def render_tabular(doc, inner):
    inner = inner.replace(r"\toprule", "").replace(r"\midrule", "@@MID@@").replace(r"\bottomrule", "")
    all_rows, head, rest = [], [], []
    for r in inner.split(r"\\"):
        r = r.strip()
        if not r:
            continue
        if "@@MID@@" in r:
            head = [x.strip() for x in r.split("@@MID@@")[0].split("&") if x.strip()]
            tail = r.split("@@MID@@")[1]
            if tail.strip():
                rest.append([x.strip() for x in tail.split("&")])
        elif head:
            rest.append([x.strip() for x in r.split("&")])
        else:
            all_rows.append([x.strip() for x in r.split("&")])
    rows = all_rows + rest
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = row[ci] if ci < len(row) else ""
            p = t.cell(ri, ci).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            emit_inline(p, cell, size=Pt(9))
            if head and ri == 0:
                for r in p.runs:
                    r.bold = True


def build():
    tex = open(SRC, encoding="utf-8").read()
    global LABELS
    blocks = parse_body(tex)
    LABELS, ENVNUM = number_labels(blocks)
    title = re.search(r"\\title\{(.*?)\}\s*\n", tex, re.S)
    author = re.search(r"\\author\{(.*?)\}\s*\n", tex, re.S)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Cambria Math"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), "宋体")
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.12
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.3)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title.group(1).strip() if title else "考拉兹截断：完整证明与开放问题")
    r.bold = True; r.font.size = Pt(16); r.font.name = "Cambria Math"; eastasia(r, "黑体")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(author.group(1).strip() if author else "")
    r.font.size = Pt(10.5); eastasia(r)
    import datetime
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(datetime.date.today().strftime("%Y 年 %m 月 %d 日"))
    r.font.size = Pt(10.5); eastasia(r)

    bib_n = 0
    for idx, b in enumerate(blocks):
        if b[0] == "sec":
            _, lvl, txt = b
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
            emit_inline(p, txt, bold=True,
                        size=Pt(14 if lvl == 1 else (12 if lvl == 2 else 11)), cjk="黑体")
        elif b[0] == "bib":
            bib_n += 1
            txt = re.sub(r"\\bibitem\{[^}]*\}", "", b[1]).strip()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Pt(21); p.paragraph_format.first_line_indent = Pt(-21)
            emit_inline(p, "[%d] %s" % (bib_n, txt), size=Pt(10))
        elif b[0] == "p":
            add_par(doc, b[1])
        elif b[0] == "env":
            _, env, opt, inner = b
            if env == "abstract":
                p = doc.add_paragraph(); r = p.add_run("摘要")
                r.bold = True; r.font.size = Pt(12); eastasia(r, "黑体")
                render_segments(doc, blocks_of(inner), size=Pt(10))
            elif env == "proof":
                segs = blocks_of(inner)
                first = True
                for kind, val in segs:
                    if kind == "m":
                        render_display(doc, val); continue
                    if kind == "t":
                        if not val.strip():
                            continue
                        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        p.paragraph_format.first_line_indent = Pt(21)
                        if first:
                            r = p.add_run("证明  "); r.bold = True; r.italic = True
                            r.font.name = "Cambria Math"; eastasia(r, "楷体")
                            first = False
                        emit_inline(p, val)
                    else:
                        style = "List Bullet" if kind == "ul" else "List Number"
                        for it in val:
                            p = doc.add_paragraph(style=style)
                            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            emit_inline(p, it)
                            if first and style == "List Bullet":
                                first = False
                continue
            elif env in NUMBERED:
                lab = re.search(r"\\label\{([^}]+)\}", inner)
                num = ENVNUM.get(idx, "?")
                inner_clean = re.sub(r"\\label\{[^}]*\}", "", inner)
                title_txt = "%s %s" % (NAMES[env], num)
                if opt:
                    title_txt += "（%s）" % flat(opt)
                segs = blocks_of(inner_clean)
                first = True
                for kind, val in segs:
                    if kind == "m":
                        if first:
                            add_par(doc, title_txt, bold=False, indent=True)
                            first = False
                        render_display(doc, val); continue
                    if kind == "t":
                        if not val.strip():
                            continue
                        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        p.paragraph_format.first_line_indent = Pt(21)
                        if first:
                            r = p.add_run(title_txt + "  ")
                            r.bold = True
                            if env == "remark":
                                r.italic = True; r.bold = False
                            r.font.name = "Cambria Math"
                            eastasia(r, "楷体" if env == "remark" else "黑体")
                            first = False
                        emit_inline(p, val)
                    else:
                        style = "List Bullet" if kind == "ul" else "List Number"
                        for it in val:
                            p = doc.add_paragraph(style=style)
                            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                            emit_inline(p, it)
                if first:
                    add_par(doc, title_txt)
                continue
            elif env in ("itemize", "enumerate"):
                style = "List Bullet" if env == "itemize" else "List Number"
                for it in [x.strip() for x in re.split(r"\\item\b", inner) if x.strip()]:
                    p = doc.add_paragraph(style=style)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    emit_inline(p, it)
            elif env == "center":
                mt = re.search(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}", inner, re.S)
                if mt:
                    render_tabular(doc, mt.group(1))
                else:
                    render_segments(doc, blocks_of(inner))
            elif env == "thebibliography":
                bn = 0
                for it in [x.strip() for x in re.split(r"\\bibitem\{[^}]*\}", inner) if x.strip()]:
                    bn += 1
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.left_indent = Pt(21)
                    p.paragraph_format.first_line_indent = Pt(-21)
                    emit_inline(p, "[%d] %s" % (bn, it), size=Pt(10))
            elif env == "display":
                render_display(doc, inner)
            else:
                render_segments(doc, blocks_of(inner))
    doc.save(DST)
    print("已生成:", DST)
    print("标签编号:", ", ".join("%s=%s" % kv for kv in sorted(LABELS.items(), key=lambda x: x[1])))
    if WARN:
        import collections
        print("未识别的命令（会被丢弃）:", dict(collections.Counter(WARN)))


if __name__ == "__main__":
    build()
